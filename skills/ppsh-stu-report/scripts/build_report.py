#!/usr/bin/env python3
"""依一份 JSON 內容，產生學習歷程報告 Word 檔（需要 python-docx）。

用法：python build_report.py <report.json> [--project <專案資料夾>]

JSON 欄位（全部文字由學生確認後才填入）：
{
  "title": "報告標題",
  "class_seat": "高一3班 15號",          # 只寫班級與座號，不寫姓名
  "course": "課程或活動名稱",
  "period": "2026-10-07 ~ 2026-10-21",
  "tools": "Codex、ComfyUI／imagegen",
  "summary": "成果簡述（約 100 字，供課程學習成果百字簡述使用）",
  "motivation": "學習動機與目標（1 段）",
  "process": [ {"step": "步驟標題", "prompt": "學生下的指令（可省略）", "note": "做了什麼、遇到什麼"} ],
  "results": [ {"file": "作業成果/圖片/xxx.png", "caption": "圖說"} ],
  "reflection": "學生省思（已經學生同意的版本）",
  "output": "作業成果/文件/學習歷程報告.docx"
}
圖片會嵌入；影片只列檔名與說明（Word 不嵌影片）。
"""
import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    print("缺少 python-docx（安裝名稱 python-docx）。請先請學生同意安裝。")
    sys.exit(3)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
VIDEO_EXT = {".mp4", ".mov", ".webm", ".avi", ".mkv", ".gif_video"}
FONT = "Microsoft JhengHei"


def set_font(style, size=None):
    style.font.name = FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        style.font.size = Pt(size)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return h


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--project", default=".")
    a = ap.parse_args()

    project = Path(a.project).resolve()
    spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    warnings = []

    doc = Document()
    set_font(doc.styles["Normal"], 12)
    for s in ("Heading 1", "Heading 2", "Title"):
        set_font(doc.styles[s])

    t = doc.add_heading(spec.get("title", "學習歷程報告"), 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 基本資料表
    info = [("班級座號", spec.get("class_seat")), ("課程／活動", spec.get("course")),
            ("期間", spec.get("period")), ("使用工具", spec.get("tools"))]
    info = [(k, v) for k, v in info if v]
    if info:
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for k, v in info:
            row = table.add_row().cells
            row[0].text, row[1].text = k, v
            row[0].width, row[1].width = Cm(3.5), Cm(12.5)

    if spec.get("summary"):
        add_heading(doc, "成果簡述")
        doc.add_paragraph(spec["summary"])

    if spec.get("motivation"):
        add_heading(doc, "一、學習動機與目標")
        doc.add_paragraph(spec["motivation"])

    process = spec.get("process") or []
    if process:
        add_heading(doc, "二、學習歷程")
        for i, p in enumerate(process, 1):
            add_heading(doc, f"{i}. {p.get('step', '')}", 2)
            if p.get("prompt"):
                para = doc.add_paragraph()
                para.add_run("我下的指令：").bold = True
                para.add_run(p["prompt"])
            if p.get("note"):
                doc.add_paragraph(p["note"])

    results = spec.get("results") or []
    if results:
        add_heading(doc, "三、成果展示")
        for r in results:
            f = Path(r.get("file", ""))
            f = f if f.is_absolute() else project / f
            ext = f.suffix.lower()
            if not f.exists():
                warnings.append(f"找不到檔案：{f}")
                continue
            if ext in IMAGE_EXT:
                doc.add_picture(str(f), width=Cm(12))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"（{'影片' if ext in VIDEO_EXT else '檔案'}：{f.name}，"
                                  f"請見作業成果資料夾）")
            if r.get("caption"):
                cap = doc.add_paragraph(r["caption"])
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if spec.get("reflection"):
        add_heading(doc, "四、學習省思")
        for para in str(spec["reflection"]).split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    add_heading(doc, "附：AI 使用說明")
    doc.add_paragraph(
        "本報告中的圖片或影片由 AI 生成；使用的工具與指令如上所列。"
        "報告文字由我口述重點，經 AI 協助整理，並已由我確認同意後才寫入。")

    out = Path(spec.get("output") or "作業成果/文件/學習歷程報告.docx")
    out = out if out.is_absolute() else project / out
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"已輸出：{out}")
    for w in warnings:
        print("警告：" + w)


if __name__ == "__main__":
    main()
